from __future__ import annotations

from dataclasses import asdict, dataclass
from pathlib import Path
from typing import List, Optional

from pptx import Presentation

from .generator import SLIDE_HEIGHT, SLIDE_WIDTH


@dataclass(frozen=True)
class ValidationIssue:
    slide: int
    type: str
    severity: str
    message: str


@dataclass(frozen=True)
class ValidationResult:
    valid: bool
    issues: List[ValidationIssue]

    def as_dict(self) -> dict:
        return {"valid": self.valid, "issues": [asdict(issue) for issue in self.issues]}


def validate_presentation(path: Path, expected_slides: Optional[int] = None) -> ValidationResult:
    prs = Presentation(path)
    issues: List[ValidationIssue] = []

    if prs.slide_width != SLIDE_WIDTH or prs.slide_height != SLIDE_HEIGHT:
        issues.append(
            ValidationIssue(0, "slide_dimensions", "error", "Presentation is not 16:9 widescreen.")
        )
    if expected_slides is not None and len(prs.slides) != expected_slides:
        issues.append(
            ValidationIssue(0, "slide_count", "error", f"Expected {expected_slides} slides, found {len(prs.slides)}.")
        )

    for slide_number, slide in enumerate(prs.slides, start=1):
        text = " ".join(
            shape.text.strip()
            for shape in slide.shapes
            if getattr(shape, "has_text_frame", False) and shape.text.strip()
        )
        if not text:
            issues.append(ValidationIssue(slide_number, "empty_slide", "error", "Slide contains no text."))
        if len(text.split()) > 110:
            issues.append(
                ValidationIssue(slide_number, "excessive_text", "warning", "Slide contains more than 110 words.")
            )
        for shape in slide.shapes:
            if shape.left < 0 or shape.top < 0 or shape.left + shape.width > prs.slide_width or shape.top + shape.height > prs.slide_height:
                # Full-bleed template polygons intentionally extend slightly beyond the canvas.
                if shape.shape_type != 1:
                    issues.append(
                        ValidationIssue(slide_number, "outside_bounds", "warning", "A shape extends outside the slide.")
                    )

    return ValidationResult(
        valid=not any(issue.severity == "error" for issue in issues),
        issues=issues,
    )
