from __future__ import annotations

import re
from pathlib import Path
from typing import Optional

from pptx import Presentation
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.util import Emu, Inches, Pt

from agent.models import PresentationSpec, SlideSpec, SlideType
from .style import STYLE, UniversityStyle


SLIDE_WIDTH = Emu(12_192_000)
SLIDE_HEIGHT = Emu(6_858_000)
ASSET_DIR = Path(__file__).resolve().parent.parent / "style_guide" / "assets"
TITLE_LOGO = ASSET_DIR / "logo-on-black.png"
CORNER_MARK = ASSET_DIR / "corner-mark.png"


def _set_fill(shape, color) -> None:
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()


def _set_text(
    shape,
    text: str,
    *,
    size: float,
    color,
    bold: bool = False,
    font: Optional[str] = None,
    align=PP_ALIGN.LEFT,
) -> None:
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.auto_size = MSO_AUTO_SIZE.NONE
    frame.margin_left = 0
    frame.margin_right = 0
    frame.margin_top = 0
    frame.margin_bottom = 0
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run()
    run.text = text
    run.font.name = font or STYLE.body_font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color


class PresentationGenerator:
    def __init__(self, style: UniversityStyle = STYLE) -> None:
        self.style = style

    def generate(self, spec: PresentationSpec, output_path: Path) -> Path:
        prs = Presentation()
        prs.slide_width = SLIDE_WIDTH
        prs.slide_height = SLIDE_HEIGHT
        blank = prs.slide_layouts[6]

        for number, slide_spec in enumerate(spec.slides, start=1):
            slide = prs.slides.add_slide(blank)
            self._render_slide(slide, slide_spec, number, len(spec.slides))

        output_path.parent.mkdir(parents=True, exist_ok=True)
        prs.save(output_path)
        return output_path

    def _render_slide(self, slide, spec: SlideSpec, number: int, total: int) -> None:
        if spec.type == SlideType.TITLE:
            self._title_slide(slide, spec)
        elif spec.type == SlideType.AGENDA:
            self._agenda_slide(slide, spec, number)
        elif spec.type == SlideType.SECTION:
            self._section_slide(slide, spec)
        elif spec.type == SlideType.QUOTE:
            self._quote_slide(slide, spec, number)
        elif spec.type == SlideType.CLOSING:
            self._closing_slide(slide, spec)
        else:
            self._content_slide(slide, spec, number)

    def _brand_mark(self, slide) -> None:
        if not CORNER_MARK.exists():
            raise FileNotFoundError(f"Official corner mark is missing: {CORNER_MARK}")
        slide.shapes.add_picture(
            str(CORNER_MARK),
            Inches(12.55),
            Inches(0.28),
            height=Inches(0.38),
        )

    def _footer(self, slide, number: int, color=None) -> None:
        page = slide.shapes.add_textbox(Inches(12.78), Inches(7.06), Inches(0.2), Inches(0.2))
        _set_text(page, str(number), size=8, color=self.style.black, align=PP_ALIGN.RIGHT)

    def _title_slide(self, slide, spec: SlideSpec) -> None:
        background = slide.background.fill
        background.solid()
        background.fore_color.rgb = self.style.black

        gold = slide.shapes.add_shape(
            MSO_AUTO_SHAPE_TYPE.PARALLELOGRAM,
            Inches(7.45), Inches(0), Inches(5.883333), Inches(7.5),
        )
        gold.rotation = 0
        _set_fill(gold, self.style.gold)

        title = slide.shapes.add_textbox(Inches(0.78), Inches(2.45), Inches(6.8), Inches(1.55))
        _set_text(title, spec.title.upper(), size=34, color=self.style.white, bold=True)

        if spec.subtitle:
            subtitle = slide.shapes.add_textbox(Inches(0.82), Inches(4.12), Inches(5.7), Inches(0.9))
            _set_text(subtitle, spec.subtitle, size=17, color=self.style.gold)

        if not TITLE_LOGO.exists():
            raise FileNotFoundError(f"Official title logo is missing: {TITLE_LOGO}")
        slide.shapes.add_picture(
            str(TITLE_LOGO),
            Inches(0.65),
            Inches(0.40),
            width=Inches(5.25),
        )

    def _agenda_slide(self, slide, spec: SlideSpec, number: int) -> None:
        panel = slide.shapes.add_shape(
            MSO_AUTO_SHAPE_TYPE.PARALLELOGRAM,
            Inches(-1.25), Inches(0), Inches(5.85), Inches(7.5),
        )
        _set_fill(panel, self.style.black)
        heading = slide.shapes.add_textbox(Inches(0.72), Inches(0.58), Inches(2.75), Inches(1.15))
        _set_text(heading, spec.title.upper(), size=24, color=self.style.gold, bold=True)

        y = 1.4
        for index, item in enumerate(spec.bullets, start=1):
            circle = slide.shapes.add_shape(
                MSO_AUTO_SHAPE_TYPE.OVAL,
                Inches(4.72), Inches(y), Inches(0.46), Inches(0.46),
            )
            _set_fill(circle, self.style.magenta)
            _set_text(circle, str(index), size=15, color=self.style.white, bold=True, align=PP_ALIGN.CENTER)
            text = slide.shapes.add_textbox(Inches(5.48), Inches(y - 0.02), Inches(6.4), Inches(0.58))
            _set_text(text, item, size=14, color=self.style.black, bold=True)
            y += 0.88

        self._brand_mark(slide)
        self._footer(slide, number)

    def _content_slide(self, slide, spec: SlideSpec, number: int) -> None:
        if number % 3 == 1:
            self._content_slide_gold(slide, spec, number)
            return

        panel = slide.shapes.add_shape(
            MSO_AUTO_SHAPE_TYPE.PARALLELOGRAM,
            Inches(-1.25), Inches(0), Inches(6.0), Inches(7.5),
        )
        _set_fill(panel, self.style.black)

        title = slide.shapes.add_textbox(Inches(0.72), Inches(1.85), Inches(3.05), Inches(1.25))
        _set_text(title, spec.title, size=22, color=self.style.gold, bold=True)
        if spec.subtitle:
            subtitle = slide.shapes.add_textbox(Inches(0.72), Inches(3.12), Inches(2.95), Inches(1.35))
            _set_text(subtitle, spec.subtitle, size=12, color=self.style.light_gray)

        y = 1.68
        spacing = min(1.0, 4.6 / max(len(spec.bullets), 1))
        for item in spec.bullets:
            marker = slide.shapes.add_shape(
                MSO_AUTO_SHAPE_TYPE.RIGHT_TRIANGLE,
                Inches(5.42), Inches(y + 0.1), Inches(0.13), Inches(0.13),
            )
            marker.rotation = 45
            _set_fill(marker, self.style.gold)
            body = slide.shapes.add_textbox(Inches(5.72), Inches(y), Inches(6.55), Inches(0.75))
            _set_text(body, item, size=18, color=self.style.black)
            y += spacing

        self._brand_mark(slide)
        self._footer(slide, number)

    def _content_slide_gold(self, slide, spec: SlideSpec, number: int) -> None:
        panel = slide.shapes.add_shape(
            MSO_AUTO_SHAPE_TYPE.RECTANGLE,
            Inches(0), Inches(0), Inches(4.15), Inches(7.5),
        )
        _set_fill(panel, self.style.gold)

        title = slide.shapes.add_textbox(Inches(0.72), Inches(2.1), Inches(2.95), Inches(1.25))
        _set_text(title, spec.title, size=24, color=self.style.white, bold=True)
        if spec.subtitle:
            subtitle = slide.shapes.add_textbox(Inches(0.72), Inches(3.55), Inches(2.85), Inches(1.4))
            _set_text(subtitle, spec.subtitle, size=12, color=self.style.black)

        y = 1.55
        spacing = min(1.05, 4.7 / max(len(spec.bullets), 1))
        for item in spec.bullets:
            marker = slide.shapes.add_shape(
                MSO_AUTO_SHAPE_TYPE.RIGHT_TRIANGLE,
                Inches(4.78), Inches(y + 0.1), Inches(0.14), Inches(0.14),
            )
            marker.rotation = 45
            _set_fill(marker, self.style.magenta)
            body = slide.shapes.add_textbox(Inches(5.12), Inches(y), Inches(7.25), Inches(0.78))
            _set_text(body, item, size=18, color=self.style.black)
            y += spacing

        self._brand_mark(slide)
        self._footer(slide, number, color=self.style.black)

    def _section_slide(self, slide, spec: SlideSpec) -> None:
        background = slide.background.fill
        background.solid()
        background.fore_color.rgb = self.style.gold
        title = slide.shapes.add_textbox(Inches(3.15), Inches(2.75), Inches(7.0), Inches(1.4))
        _set_text(title, spec.title.upper(), size=36, color=self.style.black, bold=True, align=PP_ALIGN.CENTER)
        if spec.subtitle:
            subtitle = slide.shapes.add_textbox(Inches(3.15), Inches(4.05), Inches(7.0), Inches(0.55))
            _set_text(subtitle, spec.subtitle, size=18, color=self.style.white, align=PP_ALIGN.CENTER)

    def _quote_slide(self, slide, spec: SlideSpec, number: int) -> None:
        quote_panel = slide.shapes.add_shape(
            MSO_AUTO_SHAPE_TYPE.PENTAGON,
            Inches(4.75), Inches(1.42), Inches(7.9), Inches(5.35),
        )
        quote_panel.rotation = 180
        _set_fill(quote_panel, self.style.gold)
        mark = slide.shapes.add_textbox(Inches(5.95), Inches(1.75), Inches(0.8), Inches(0.7))
        _set_text(mark, "“", size=56, color=self.style.magenta, bold=True)
        quote = slide.shapes.add_textbox(Inches(6.72), Inches(2.0), Inches(4.85), Inches(2.65))
        _set_text(quote, spec.quote or "", size=19, color=self.style.white, bold=True)
        if spec.attribution:
            attribution = slide.shapes.add_textbox(Inches(6.72), Inches(5.05), Inches(4.5), Inches(0.8))
            _set_text(attribution, spec.attribution, size=11, color=self.style.black)
        self._brand_mark(slide)
        self._footer(slide, number)

    def _closing_slide(self, slide, spec: SlideSpec) -> None:
        background = slide.background.fill
        background.solid()
        background.fore_color.rgb = self.style.black
        gold = slide.shapes.add_shape(
            MSO_AUTO_SHAPE_TYPE.PARALLELOGRAM,
            Inches(0), Inches(0), Inches(7.7), Inches(7.5),
        )
        _set_fill(gold, self.style.gold)
        title = slide.shapes.add_textbox(Inches(8.0), Inches(2.75), Inches(4.5), Inches(1.0))
        _set_text(title, spec.title.upper(), size=34, color=self.style.white, bold=True)
        if spec.subtitle:
            subtitle = slide.shapes.add_textbox(Inches(8.04), Inches(4.9), Inches(4.25), Inches(1.0))
            _set_text(subtitle, spec.subtitle, size=12, color=self.style.gold)


def safe_filename(title: str) -> str:
    value = re.sub(r"[^a-z0-9]+", "-", title.lower()).strip("-")
    return (value[:70] or "presentation") + ".pptx"
