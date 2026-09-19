from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from agent.models import PresentationSpec
from presentation.generator import PresentationGenerator
from presentation.validator import validate_presentation


def sample_plan() -> PresentationSpec:
    return PresentationSpec.model_validate(
        {
            "title": "Quantum Computing",
            "subtitle": "An introduction",
            "audience": "First-year university students",
            "slides": [
                {"type": "title", "title": "Quantum Computing", "subtitle": "An introduction"},
                {"type": "content", "title": "Qubits", "bullets": ["A qubit stores quantum information", "Measurement produces a classical result"]},
                {"type": "section", "title": "Applications", "subtitle": "Where quantum methods may help"},
                {"type": "content", "title": "Current limitations", "bullets": ["Noise disrupts calculations", "Error correction requires many physical qubits"]},
                {"type": "closing", "title": "Thank you"},
            ],
        }
    )


def test_generator_creates_readable_pptx(tmp_path: Path) -> None:
    output = tmp_path / "sample.pptx"
    PresentationGenerator().generate(sample_plan(), output)
    reopened = Presentation(output)
    assert len(reopened.slides) == 5
    assert any(shape.shape_type == MSO_SHAPE_TYPE.PICTURE for shape in reopened.slides[0].shapes)
    assert any(shape.shape_type == MSO_SHAPE_TYPE.PICTURE for shape in reopened.slides[1].shapes)
    assert "KNOWLEDGE FOUNDATION\nReutlingen University" not in [
        shape.text
        for shape in reopened.slides[0].shapes
        if getattr(shape, "has_text_frame", False)
    ]
    assert not any(
        "Weiterbildung an der Hochschule Reutlingen" in shape.text
        for slide in reopened.slides
        for shape in slide.shapes
        if getattr(shape, "has_text_frame", False)
    )
    result = validate_presentation(output, expected_slides=5)
    assert result.valid
